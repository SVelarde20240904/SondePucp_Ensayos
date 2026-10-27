"""
Genera db/seed.sql a partir del Excel "Matriz_Figuras_SDP_v2_6.xlsx".
Uso:  python generar_seed.py ruta/al/Matriz_Figuras_SDP.xlsx  > ../seed.sql
Requiere: pip install openpyxl
"""
import sys, datetime as dt, unicodedata
import openpyxl

XLSX = sys.argv[1]
wb = openpyxl.load_workbook(XLSX, data_only=True)
out = []
w = out.append

def q(v):
    if v is None or v == "":
        return "NULL"
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, (int, float)):
        return str(int(v))
    if isinstance(v, dt.datetime):
        return "'" + v.strftime("%Y-%m-%d %H:%M:%S") + "'"
    if isinstance(v, dt.date):
        return "'" + v.strftime("%Y-%m-%d") + "'"
    return "'" + str(v).strip().replace("\\", "\\\\").replace("'", "''") + "'"

def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return " ".join(s.lower().split())

w("-- =====================================================================")
w("--  Datos iniciales migrados desde el Excel (generado por generar_seed.py)")
w(f"--  Fuente: {XLSX.split('/')[-1]} · generado {dt.date.today()}")
w("-- =====================================================================")
w("USE sdp_ensayos;")
w("SET NAMES utf8mb4;")
w("START TRANSACTION;\n")

# ---------------- Catálogos ----------------
w("-- Roles del sistema")
w("INSERT INTO rol_sistema (id_rol, codigo, nombre) VALUES (1,'MIEMBRO','Miembro'),(2,'EQUIPO_ENSAYOS','Equipo de ensayos'),(3,'ADMIN','Admin');\n")
perms = [
 ("CONFIRMAR_ASISTENCIA","Confirmar asistencia","MI_ESPACIO",(1,1,1)),
 ("AUTOEVALUARSE","Autoevaluarse (mensual y del ensayo)","MI_ESPACIO",(1,1,1)),
 ("VER_HISTORIAL_PROPIO","Ver su historial y perfil","MI_ESPACIO",(1,1,1)),
 ("GESTIONAR_ENSAYOS","Crear y editar ensayos","ENSAYOS",(0,1,1)),
 ("VER_CRONOGRAMA","Ver cronograma sugerido","ENSAYOS",(0,1,1)),
 ("TOMAR_ASISTENCIA","Tomar asistencia","ENSAYOS",(0,1,1)),
 ("REGISTRAR_FIGURAS_ENSENADAS","Registrar figuras enseñadas","ENSAYOS",(0,1,1)),
 ("IMPRIMIR_REPORTES","Imprimir reportes de ensayo","ENSAYOS",(0,1,1)),
 ("GESTIONAR_MIEMBROS","Gestionar miembros","ADMINISTRACION",(0,0,1)),
 ("GESTIONAR_FIGURAS","Gestionar mano de figuras","ADMINISTRACION",(0,0,1)),
 ("ASIGNAR_RESPONSABILIDADES","Asignar responsabilidades","ADMINISTRACION",(0,1,1)),
 ("GESTIONAR_HISTORICO","Ver y corregir histórico","ADMINISTRACION",(0,0,1)),
 ("ASIGNAR_ROLES","Asignar roles","ADMINISTRACION",(0,0,1)),
]
w("-- Permisos (pantalla 19)")
w("INSERT INTO permiso (id_permiso, codigo, descripcion, seccion) VALUES")
w(",\n".join(f"  ({i+1},{q(c)},{q(d)},{q(s)})" for i,(c,d,s,_) in enumerate(perms)) + ";")
rp = [f"({r+1},{i+1})" for i,(_,_,_,v) in enumerate(perms) for r in range(3) if v[r]]
w("INSERT INTO rol_permiso (id_rol, id_permiso) VALUES " + ",".join(rp) + ";\n")

# ---------------- Reglas de conteo ----------------
m = wb["MIEMBROS"]
reglas = {}
for r in range(2, 5):
    k, v = m.cell(r, 10).value, m.cell(r, 11).value
    if k: reglas[norm(k).upper()] = (str(v).upper() == "SI")
reglas["EQUIPO_ENSAYOS"] = str(m["K7"].value).upper() == "SI"
w("-- Quién cuenta para el color (antes MIEMBROS!J2:K4 y K7)")
w("INSERT INTO regla_conteo_color (clave, cuenta) VALUES " + ",".join(f"({q(k)},{q(v)})" for k,v in reglas.items()) + ";\n")

# ---------------- Miembros ----------------
ROL = {"lider": "LEADER", "seguidora": "FOLLOWER", "leader": "LEADER", "follower": "FOLLOWER"}
miembros = []  # (id, nombre)
w("-- Miembros (hoja MIEMBROS)")
w("INSERT INTO miembro (id_miembro, nombre, nombre_completo, rol_baile, frecuencia, estado, fecha_ingreso, notas) VALUES")
rows = []
equipo = []
for r in m.iter_rows(min_row=2, max_col=8, values_only=True):
    if not r[0]: continue
    idm = len(miembros) + 1
    miembros.append((idm, r[0].strip()))
    rol = ROL.get(norm(r[2])) if r[2] else None
    fr = norm(r[3]).upper() if r[3] else None
    est = norm(r[5]).upper() if r[5] else "ACTIVO"
    rows.append(f"  ({idm},{q(r[0])},{q(r[1] or r[0])},{q(rol)},{q(fr)},{q(est)},{q(r[6])},{q(r[7])})")
    if str(r[4]).upper() == "SI": equipo.append(idm)
w(",\n".join(rows) + ";\n")
ID = {norm(n): i for i, n in miembros}

def mid(name):
    k = norm(name)
    if k in ID: return ID[k]
    raise KeyError(name)

w("-- Roles extra (el rol MIEMBRO lo pone el trigger trg_miembro_ai): equipo de ensayos y admin")
mr = [f"({i},2)" for i in equipo] + [f"({mid('Sergio Velarde')},3)"]
w("INSERT INTO miembro_rol (id_miembro, id_rol) VALUES " + ",".join(mr) + ";\n")

# ---------------- Figuras ----------------
f = wb["FIGURAS"]
figs = []
for r in f.iter_rows(min_row=2, max_col=2, values_only=True):
    if r[0]: figs.append((len(figs)+1, r[0].strip(), int(r[1]) if r[1] else len(figs)+1))
FID = {norm(n): i for i, n, _ in figs}
w("-- Figuras (hoja FIGURAS)")
w("INSERT INTO figura (id_figura, nombre, orden) VALUES " + ",\n  ".join(f"({i},{q(n)},{o})" for i,n,o in figs) + ";\n")

# ---------------- Periodo de autoevaluación ----------------
w("-- Periodo mensual en que se hizo la primera autoevaluación")
w("INSERT INTO periodo_autoevaluacion (id_periodo, anio, mes, fecha_apertura, fecha_cierre) VALUES (1,2026,10,'2026-10-22 00:00:00','2026-10-31 23:59:59');\n")

# ---------------- Matriz + Histórico ----------------
mat = wb["MATRIZ"]
hdr = [c.value for c in mat[1]]
current = {}
for r in range(2, mat.max_row + 1):
    fig = mat.cell(r, 1).value
    if not fig: continue
    for c in range(2, len(hdr) + 1):
        name = hdr[c-1]
        if not name or norm(name) not in ID: continue
        v = mat.cell(r, c).value
        if v in ("ROJO","AMARILLO","VERDE","AZUL"):
            current[(mid(name), FID[norm(fig)])] = v

hist = wb["HISTORICO"]
hrows = [r for r in hist.iter_rows(min_row=2, max_col=5, values_only=True) if r[0] and r[1]]
last = {}
hist_sql = []
seen_prev = {}
for fecha, name, fig, color, fuente in sorted(hrows, key=lambda r: r[0]):
    key = (mid(name), FID[norm(fig)])
    prev = seen_prev.get(key)
    src = {"mensual": "MENSUAL", "semanal": "ENSAYO", "correccion": "CORRECCION", "carga_inicial": "CARGA_INICIAL"}.get(norm(fuente).replace(" ", "_"), "MENSUAL")
    per = "1" if src == "MENSUAL" else "NULL"
    hist_sql.append(f"({key[0]},{key[1]},{q(prev)},{q(color)},{q(src)},1,{per},{key[0]},{q(fecha)})")
    seen_prev[key] = color
    last[key] = (color, fecha)

# Celdas de la matriz sin historial -> carga inicial
carga = dt.datetime(2026, 10, 22)
mismatch = 0
for key, color in current.items():
    if key not in last:
        hist_sql.append(f"({key[0]},{key[1]},NULL,{q(color)},'CARGA_INICIAL',1,NULL,NULL,{q(carga)})")
    elif last[key][0] != color:
        mismatch += 1
w(f"-- Matriz actual: {len(current)} celdas con color")
w("INSERT INTO evaluacion (id_miembro, id_figura, color, fecha_actualizacion) VALUES")
w(",\n".join(f"  ({k[0]},{k[1]},{q(c)},{q(last[k][1] if k in last else carga)})" for k, c in sorted(current.items())) + ";\n")
w(f"-- Histórico: {len(hrows)} filas del Excel + {len(hist_sql)-len(hrows)} de carga inicial (celdas sin historial)")
if mismatch: w(f"-- OJO: {mismatch} celdas de la matriz no coinciden con su último registro del histórico (se respeta la matriz)")
w("INSERT INTO historial_evaluacion (id_miembro, id_figura, color_anterior, color_nuevo, fuente, aplicado, id_periodo, registrado_por, fecha) VALUES")
w(",\n".join("  " + s for s in hist_sql) + ";\n")

# ---------------- Ensayos + asistencia (hoja ASISTENCIA) ----------------
ALIAS_ASIS = {"alessia": "Alessia Leon", "alex": "Alexander Gil", "alvaro": "Alvaro Camposano", "andre": "André La Hoz",
    "angela": "Angela Yanac", "ariana mercedes": "Ariana David", "cielo": "Cielo Camayo", "daniela lisseth": "Daniela Llano",
    "deyvid": "Deyvid Usco", "erich vercelli": "Erick Vercelli", "esteban": "Esteban Gutiérrez", "enmanuel (ilo)": "Enmanuel Vega",
    "jose perez": "José Pérez", "mafer alvarez": "Maria Fernanda Alvarez", "maricielo": "Maricielo Mamani", "mauricio": "Mauricio Halanoca"}
a = wb["ASISTENCIA"]
fechas = [c.value for c in a[1][1:] if isinstance(c.value, dt.datetime)]
w("-- Ensayos históricos (columnas de la hoja ASISTENCIA)")
ens = []
for i, d in enumerate(fechas, 1):
    sab = d.weekday() == 5
    hi, hf = ("14:00:00", "16:00:00") if sab else ("20:00:00", "22:00:00")
    ens.append(f"  ({i},'GENERAL',{q(d.date())},'{hi}','{hf}','Explanada del polideportivo','CERRADO',{mid('Sergio Velarde')})")
w("INSERT INTO ensayo (id_ensayo, tipo, fecha, hora_inicio, hora_fin, lugar, estado, creado_por) VALUES")
w(",\n".join(ens) + ";\n")
EID = {d.date(): i for i, d in enumerate(fechas, 1)}

CONF = {"AP": "SI", "TJ": "SI", "TI": "SI", "FJ": "SI", "FI": "SI", "AN": "NO"}
asis_rows = []
sin_mapear = []
for r in a.iter_rows(min_row=2, values_only=True):
    name = r[0]
    if not name or norm(name) in ("leyenda",) or (isinstance(name, str) and norm(name) in ("x","ap","an","fi","fj","tj","ti")):
        if name and norm(name) == "leyenda": break
        continue
    k = norm(name)
    target = ALIAS_ASIS.get(k, name)
    try: idm = mid(target)
    except KeyError:
        sin_mapear.append(name); continue
    for j, d in enumerate(fechas):
        v = r[j+1]
        if v is None: continue
        v = str(v).strip().upper()
        if v == "X":
            asis_rows.append(f"({EID[d.date()]},{idm},'SIN_RESPUESTA',NULL)")
        elif v in CONF:
            asis_rows.append(f"({EID[d.date()]},{idm},'{CONF[v]}','{v}')")
w("-- Asistencia: 'x' = no marcó la encuesta -> SIN_RESPUESTA (ver README si en realidad significaba 'marcó que no')")
if sin_mapear: w("-- Nombres de ASISTENCIA sin miembro: " + ", ".join(sin_mapear))
w("INSERT INTO asistencia (id_ensayo, id_miembro, confirmacion, estado_real) VALUES")
w(",\n".join("  " + s for s in asis_rows) + ";\n")

# ---------------- Responsabilidades ----------------
ALIAS_RESP = {"antony": "Antony Palacios", "sergio": "Sergio Velarde", "pierina": "Pierina Perales", "claudia": "Claudia Alonso", "clau": "Claudia Alonso",
    "mafer": "Maria Fernanda Alvarez", "deyvid": "Deyvid Usco", "arturo": "Arturo Custodio", "gaby": "Gabriela Arrieta", "heidy": "Heidy Sanchez",
    "josefh": "Josefh", "maricielo": "Maricielo Mamani", "pedro": "Pedro Chero"}
TIPOS = [("FLYER Y SPEECH", "Flyer y speech", 1), ("REGISTRAR ASISTENCIA Y FIGURAS", "Registrar asistencia y figuras", 1),
         ("RECOPILAR AUTOEVALUACIÓN", "Recopilar autoevaluación", 0), ("ACTUALIZAR LA BD CON LA ATUOEVALUACIÓN", "Actualizar la BD con la autoevaluación", 0),
         ("CRONOGRAMA SEMANAL", "Cronograma semanal", 0), ("CREACIÓN Y ENVIO DEL FORMULRIO", "Creación y envío del formulario", 0),
         ("ACTUALIZACIÓN COMPLETA DE LA BD", "Actualización completa de la BD", 0)]
w("-- Tipos de encargo (activo=0: los automatizó la web)")
tipos_sql = [f"({i+1},{q(n)},{a_})" for i,(_,n,a_) in enumerate(TIPOS)] + [f"({len(TIPOS)+1},'Compartir cronograma',1)", f"({len(TIPOS)+2},'Tomar asistencia',1)", f"({len(TIPOS)+3},'Registrar figuras enseñadas',1)"]
w("INSERT INTO tipo_encargo (id_tipo, nombre, activo) VALUES " + ",".join(tipos_sql) + ";\n")

rs = wb["RESPONSABILIDADES DE LA SEMANA"]
grid = [[c.value for c in row] for row in rs.iter_rows(min_row=1, max_row=rs.max_row, max_col=7)]
semanas = [v.date() for v in grid[0][1:] if isinstance(v, dt.datetime)]
# estado por persona/semana (para NCJ y x)
estado_pers = {}
start = next(i for i, row in enumerate(grid) if row[0] and "MIEMBROS VS" in str(row[0]).upper())
for row in grid[start+1:]:
    if not row[0] or str(row[0]).upper() == "LEYENDA": break
    for j, s in enumerate(semanas):
        if row[j+1]: estado_pers[(norm(row[0]), s)] = str(row[j+1]).strip().upper()

def resp_id(txt):
    first = str(txt).replace(" y ", "+").split("+")[0].strip()
    return mid(ALIAS_RESP.get(norm(first), first))

enc_sql = []
for ti, (titulo, _, _) in enumerate(TIPOS, 1):
    idx = next(i for i, row in enumerate(grid) if row[0] and norm(row[0]) == norm(titulo))
    notas, resp, fec, hor = grid[idx], grid[idx+1], grid[idx+2], grid[idx+3]
    for j, s in enumerate(semanas):
        nota = notas[j+1]; r_ = resp[j+1]
        if nota and "NO FUE NECESARIO" in str(nota).upper(): continue
        if not r_: continue
        rid = resp_id(r_)
        f_ = fec[j+1]; h_ = hor[j+1]
        fecha = None
        if isinstance(f_, dt.datetime):
            fecha = dt.datetime.combine(f_.date(), h_ if isinstance(h_, dt.time) else dt.time(0, 0))
        cumple = rid if fecha else None
        if nota and str(nota).lower().startswith("cubri"):
            cumple = resp_id(str(nota).split()[-1]); fecha = fecha or None
        alias = next((k for k, v in ALIAS_RESP.items() if mid(v) == rid), "")
        est = estado_pers.get((alias, s), "")
        just = 1 if (fecha is None and est == "NCJ") else 0
        ens_id = EID.get(s, None)
        enc_sql.append(f"({ti},{ens_id if ens_id else 'NULL'},{q(s)},{rid},NULL,{cumple if cumple else 'NULL'},{q(fecha)},0,{just},{q(str(nota) if nota and not str(nota).lower().startswith('cubri') else None)})")
w("-- Encargos históricos (hoja RESPONSABILIDADES DE LA SEMANA). Plazo desconocido -> NULL")
w("INSERT INTO encargo (id_tipo, id_ensayo, semana, responsable, plazo, cumplido_por, fecha_cumplimiento, exonerado, justificada, motivo_justificacion) VALUES")
w(",\n".join("  " + s for s in enc_sql) + ";\n")

w("COMMIT;")
print("\n".join(out))
print(f"-- resumen: {len(miembros)} miembros, {len(figs)} figuras, {len(current)} evaluaciones, {len(hist_sql)} historial, {len(fechas)} ensayos, {len(asis_rows)} asistencias, {len(enc_sql)} encargos", file=sys.stderr)
if sin_mapear: print("-- SIN MAPEAR:", sin_mapear, file=sys.stderr)
