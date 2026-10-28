#!/usr/bin/env python3
"""
Importa a la BD lo que FALTA desde la matriz original del grupo ("Matriz_Figuras_SDP_1.xlsx",
hojas MATRIZ DE FIGURAS / ASISTENCIA / RESPONSABILIDADES DE LA SEMANA).

Regla: lo que ya está en la BD (cargado desde la v2.6 o hecho en el piloto) MANDA. Este script solo agrega:
  · colores de figuras que la persona todavía no tiene
  · 4 figuras que faltaban (Peruana, La trenza, Adios con la Suegra, 8 de Mujeres). "Paseos" se ignora:
    en la web se dividió en Paseo Cerrado / Abierto / hacia arriba.
  · ensayos (fechas) que no existen, con su asistencia; y asistencias que falten en fechas que sí existen
  · encargos de semanas que no están cargadas

Genera SQL que busca por NOMBRE y FECHA (no por id), así que sirve para la BD del piloto sin tocar cuentas.
Se puede correr dos veces: no duplica nada.

Uso:  python3 importar_matriz_original.py Matriz_Figuras_SDP_1.xlsx > actualizacion.sql
El .sql tiene datos reales: NO subirlo a GitHub.
"""
import sys, re, unicodedata, datetime as dt
import openpyxl

def norm(s):
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().strip().lower())

def q(v):
    if v is None: return "NULL"
    if isinstance(v, dt.datetime): return "'" + v.strftime("%Y-%m-%d %H:%M:%S") + "'"
    if isinstance(v, dt.date): return "'" + v.isoformat() + "'"
    return "'" + str(v).replace("\\", "\\\\").replace("'", "''") + "'"

# Columnas de la matriz (apodos) -> nombre en el sistema
ALIAS_MATRIZ = {"sergio": "Sergio Velarde", "antony": "Antony Palacios", "maricielo": "Maricielo Mamani", "ilo": "Enmanuel Vega",
    "paola": "Paola Nieto", "esteban": "Esteban Gutiérrez", "alessia": "Alessia Leon", "haniel": "Haniel Ghiggo", "cielo": "Cielo Camayo",
    "daniela": "Daniela Llano", "brian": "Brian Chero", "abigail": "Abigail Delgado", "andre": "André La Hoz", "adrian": "Adrian Lopez",
    "ximena": "Ximena Robles", "beteta": "Alexander Beteta", "pierina": "Pierina Perales", "mafer": "Maria Fernanda Alvarez",
    "mimi": "Miranda Cueva", "arturo": "Arturo Custodio", "deyvid": "Deyvid Usco", "jose": "José Pérez", "claudia": "Claudia Alonso",
    "natalia": "Natalia Vidal", "diana": "Diana Levano"}
# Nombres de la hoja ASISTENCIA que no coinciden tal cual
ALIAS_ASIS = {"alessia": "Alessia Leon", "alex": "Alexander Gil", "alvaro": "Alvaro Camposano", "andre": "André La Hoz",
    "angela": "Angela Yanac", "ariana mercedes": "Ariana David", "cielo": "Cielo Camayo", "daniela lisseth": "Daniela Llano",
    "deyvid": "Deyvid Usco", "erich vercelli": "Erick Vercelli", "esteban": "Esteban Gutiérrez", "enmanuel (ilo)": "Enmanuel Vega",
    "jose perez": "José Pérez", "mafer alvarez": "Maria Fernanda Alvarez", "maricielo": "Maricielo Mamani", "mauricio": "Mauricio Halanoca",
    "gabriela old": "Gabriela Arrieta"}
ALIAS_RESP = {"antony": "Antony Palacios", "sergio": "Sergio Velarde", "pierina": "Pierina Perales", "claudia": "Claudia Alonso",
    "clau": "Claudia Alonso", "mafer": "Maria Fernanda Alvarez", "deyvid": "Deyvid Usco", "arturo": "Arturo Custodio",
    "gaby": "Gabriela Arrieta", "heidy": "Heidy Sanchez", "josefh": "Josefh", "maricielo": "Maricielo Mamani", "pedro": "Pedro Chero"}
FIGURA_ALIAS = {"botala para atras": "Botala por atrás"}
FIGURAS_NUEVAS = ["Peruana", "La trenza", "Adios con la Suegra", "8 de Mujeres"]
FIGURAS_IGNORADAS = {"paseos"}
COLORES = {"AZUL", "VERDE", "AMARILLO", "ROJO"}
CONF = {"AP": "SI", "TJ": "SI", "TI": "SI", "FJ": "SI", "FI": "SI", "AN": "NO", "ANP": "NO"}
TIPOS = ["FLYER Y SPEECH", "REGISTRAR ASISTENCIA Y FIGURAS", "RECOPILAR AUTOEVALUACIÓN", "ACTUALIZAR LA BD CON LA ATUOEVALUACIÓN",
         "CRONOGRAMA SEMANAL", "CREACIÓN Y ENVIO DEL FORMULRIO", "ACTUALIZACIÓN COMPLETA DE LA BD"]   # id_tipo = posición + 1 (como el seed)

wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
out, avisos = [], []
w = out.append
mid = lambda nombre: f"(SELECT id_miembro FROM miembro WHERE nombre={q(nombre)})"

w("-- Actualización desde la matriz original (Excel v1). Lo que ya existe en la BD NO se modifica.")
w("SET NAMES utf8mb4;")
w("START TRANSACTION;\n")

# ---------------- 1. Figuras nuevas ----------------
w("-- 1. Figuras que faltaban (van al final de la mano)")
for f in FIGURAS_NUEVAS:
    w(f"INSERT INTO figura (nombre) SELECT {q(f)} FROM DUAL WHERE NOT EXISTS (SELECT 1 FROM figura WHERE nombre={q(f)});")

# ---------------- 2. Colores de la matriz ----------------
m = wb["MATRIZ DE FIGURAS"]
hdr = [c.value for c in m[1]]
cols = {}
for i, h in enumerate(hdr):
    if h is None or i == 0: continue
    if norm(h) in ALIAS_MATRIZ: cols[i] = ALIAS_MATRIZ[norm(h)]
    elif norm(h) == "figura": break          # empieza el bloque de cálculos de la hoja
celdas = []
for row in m.iter_rows(min_row=2, values_only=True):
    fig = row[0]
    if not fig: continue
    if norm(fig) == "figura": break          # segundo bloque (números)
    if norm(fig) in FIGURAS_IGNORADAS: continue
    fig = FIGURA_ALIAS.get(norm(fig), fig)
    for i, mem in cols.items():
        v = row[i]
        v = v.strip().upper() if isinstance(v, str) else v
        if v in COLORES: celdas.append(f"({q(mem)},{q(fig)},'{v}')")
w("\n-- 2. Colores: solo se agregan los que la persona todavía no tiene")
w("DROP TEMPORARY TABLE IF EXISTS tmp_eval;")
w("CREATE TEMPORARY TABLE tmp_eval (miembro VARCHAR(60), figura VARCHAR(60), color VARCHAR(10)) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci;")
w("INSERT INTO tmp_eval VALUES\n  " + ",\n  ".join(celdas) + ";")
w("""INSERT INTO historial_evaluacion (id_miembro, id_figura, color_anterior, color_nuevo, fuente, aplicado, motivo)
SELECT m.id_miembro, f.id_figura, NULL, t.color, 'CARGA_INICIAL', 1, 'Matriz original del grupo (Excel v1)'
FROM tmp_eval t JOIN miembro m ON m.nombre = t.miembro JOIN figura f ON f.nombre = t.figura
LEFT JOIN evaluacion e ON e.id_miembro = m.id_miembro AND e.id_figura = f.id_figura
WHERE e.id_miembro IS NULL;""")
w("""INSERT IGNORE INTO evaluacion (id_miembro, id_figura, color, fecha_actualizacion)
SELECT m.id_miembro, f.id_figura, t.color, NOW()
FROM tmp_eval t JOIN miembro m ON m.nombre = t.miembro JOIN figura f ON f.nombre = t.figura;""")

# ---------------- 3. Ensayos + asistencia ----------------
a = wb["ASISTENCIA"]
fechas = [c.value for c in a[1][1:] if isinstance(c.value, dt.datetime)]
w("\n-- 3. Ensayos que no existen (por fecha) y asistencias que falten")
for d in fechas:
    hi, hf = ("14:00:00", "16:00:00") if d.weekday() == 5 else ("20:00:00", "22:00:00")
    w(f"INSERT INTO ensayo (tipo, fecha, hora_inicio, hora_fin, lugar, estado, creado_por) SELECT 'GENERAL', {q(d.date())}, '{hi}', '{hf}', "
      f"'Explanada del polideportivo', 'CERRADO', {mid('Sergio Velarde')} FROM DUAL WHERE NOT EXISTS (SELECT 1 FROM ensayo WHERE fecha={q(d.date())});")
asis = []
for r in a.iter_rows(min_row=2, values_only=True):
    name = r[0]
    if not name: continue
    if norm(name) == "leyenda": break
    target = ALIAS_ASIS.get(norm(name), str(name).strip())
    for j, d in enumerate(fechas):
        v = r[j + 1]
        if v is None: continue
        v = str(v).strip().upper()
        if v == "X": asis.append(f"({q(target)},{q(d.date())},'SIN_RESPUESTA',NULL)")
        elif v in CONF: asis.append(f"({q(target)},{q(d.date())},'{CONF[v]}','{'AN' if v == 'ANP' else v}')")
w("DROP TEMPORARY TABLE IF EXISTS tmp_asis;")
w("CREATE TEMPORARY TABLE tmp_asis (miembro VARCHAR(60), fecha DATE, confirmacion VARCHAR(15), estado_real VARCHAR(3)) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci;")
w("INSERT INTO tmp_asis VALUES\n  " + ",\n  ".join(asis) + ";")
w("""INSERT IGNORE INTO asistencia (id_ensayo, id_miembro, confirmacion, estado_real)
SELECT (SELECT MIN(e.id_ensayo) FROM ensayo e WHERE e.fecha = t.fecha), m.id_miembro, t.confirmacion, t.estado_real
FROM tmp_asis t JOIN miembro m ON m.nombre = t.miembro;""")

# ---------------- 4. Encargos de semanas nuevas ----------------
rs = wb["RESPONSABILIDADES DE LA SEMANA"]
grid = [[c.value for c in row] for row in rs.iter_rows(min_row=1, max_row=rs.max_row)]
ncols = len(grid[0])
semanas = [(j, v.date()) for j, v in enumerate(grid[0]) if j > 0 and isinstance(v, dt.datetime)]
estado_pers = {}
start = next(i for i, row in enumerate(grid) if row[0] and "MIEMBROS VS" in str(row[0]).upper())
for row in grid[start + 1:]:
    if not row[0] or str(row[0]).upper() == "LEYENDA": break
    for j, s in semanas:
        if row[j]: estado_pers[(norm(row[0]), s)] = str(row[j]).strip().upper()

def resp(txt):
    first = str(txt).replace(" y ", "+").split("+")[0].strip()
    return ALIAS_RESP.get(norm(first), first)

def es_cobertura(nota):
    return nota is not None and norm(nota).startswith(("cubri", "cubiro"))

w("\n-- 4. Encargos de semanas que todavía no están cargadas")
for ti, titulo in enumerate(TIPOS, 1):
    idx = next(i for i, row in enumerate(grid) if row[0] and norm(row[0]) == norm(titulo))
    notas, resps, fec, hor = grid[idx], grid[idx + 1], grid[idx + 2], grid[idx + 3]
    for j, s in semanas:
        nota, r_ = notas[j], resps[j]
        if nota and "NO FUE NECESARIO" in str(nota).upper(): continue
        if not r_: continue
        responsable = resp(r_)
        f_, h_ = fec[j], hor[j]
        fecha = None
        if isinstance(f_, dt.datetime):
            fd = f_.date()
            if fd - s > dt.timedelta(days=20):      # "27/07" en la semana del 27/06: error de tipeo del mes
                fd2 = fd.replace(month=fd.month - 1)
                avisos.append(f"{titulo} semana {s}: fecha {fd} corregida a {fd2} (parece error de tipeo)")
                fd = fd2
            hora = h_.time() if isinstance(h_, dt.datetime) else (h_ if isinstance(h_, dt.time) else dt.time(0, 0))
            fecha = dt.datetime.combine(fd, hora)
        cumple = responsable if fecha else None
        if es_cobertura(nota): cumple = resp(str(nota).split()[-1])
        est = estado_pers.get((next((k for k, v in ALIAS_RESP.items() if v == responsable), ""), s), "")
        just = 1 if (fecha is None and est == "NCJ") else 0
        motivo = None if (nota is None or es_cobertura(nota)) else str(nota)
        w(f"INSERT INTO encargo (id_tipo, id_ensayo, semana, responsable, plazo, cumplido_por, fecha_cumplimiento, exonerado, justificada, motivo_justificacion) "
          f"SELECT {ti}, (SELECT MIN(id_ensayo) FROM ensayo WHERE fecha={q(s)}), {q(s)}, {mid(responsable)}, NULL, "
          f"{mid(cumple) if cumple else 'NULL'}, {q(fecha)}, 0, {just}, {q(motivo)} FROM DUAL "
          f"WHERE NOT EXISTS (SELECT 1 FROM encargo WHERE id_tipo={ti} AND semana={q(s)});")

w("\nDROP TEMPORARY TABLE IF EXISTS tmp_eval;\nDROP TEMPORARY TABLE IF EXISTS tmp_asis;")
w("COMMIT;")
print("\n".join(out))
print(f"-- {len(celdas)} colores en la matriz, {len(fechas)} fechas y {len(asis)} marcas de asistencia, {len(semanas)} semanas de encargos", file=sys.stderr)
for a_ in avisos: print("-- AVISO: " + a_, file=sys.stderr)
